import os
import datetime
import pandas as pd

import streamlit as st
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


def format_tgl_indo(date_obj, is_draft=False, include_year=True):
    months = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
              "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    if is_draft:
        return f"{months[date_obj.month]} {date_obj.year}" if include_year else f"{months[date_obj.month]}"
    else:
        return f"{date_obj.day} {months[date_obj.month]} {date_obj.year}" if include_year else f"{date_obj.day} {months[date_obj.month]}"

convert_date_to_indonesian = format_tgl_indo


def set_document_margins(doc):
    # Global font: Arial 12
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(12)

    # Page setup: A4 + margins
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.0)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)


def add_kop_surat(doc):
    # 1. Create a 2-column invisible table (3.5 cm + 12.5 cm = 16 cm text width)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.5)
    table.columns[1].width = Cm(12.5)
    # Word honours per-cell widths, so set them too
    table.cell(0, 0).width = Cm(3.5)
    table.cell(0, 1).width = Cm(12.5)

    # 2. Insert Logo in Left Cell
    cell_left = table.cell(0, 0)
    cell_left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p_logo = cell_left.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_after = Pt(0)
    p_logo.paragraph_format.line_spacing = 1.0
    try:
        r_logo = p_logo.add_run()
        # Width only, so the aspect ratio stays locked
        r_logo.add_picture("logo_bpkp.png", width=Cm(3.0))
    except Exception:
        p_logo.add_run("[LOGO BPKP MISSING]").bold = True

    # 3. Insert Text in Right Cell
    cell_right = table.cell(0, 1)

    def fmt(p):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0  # Single spacing for header

    def run(p, text, size, bold=False, italic=False):
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        return r

    p1 = cell_right.paragraphs[0]
    fmt(p1)
    run(p1, "BADAN PENGAWASAN KEUANGAN DAN PEMBANGUNAN", 12, bold=True)

    p2 = cell_right.add_paragraph()
    fmt(p2)
    run(p2, "PERWAKILAN PROVINSI PAPUA", 12, bold=True)

    p3 = cell_right.add_paragraph()
    fmt(p3)
    run(p3, "Jalan Pasifik Indah III, Pasir Dua, Jayapura, Papua 99117\nTelepon (0967) 5144006, ", 9)
    run(p3, "Email", 9, italic=True)
    run(p3, " papua@bpkp.go.id\n", 9)
    run(p3, "Website", 9, italic=True)
    run(p3, " www.bpkp.go.id/papua", 9)

    # 4. Thick bottom border attached directly to the table
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        tblBorders = OxmlElement('w:tblBorders')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '24')  # 24 = 3.0 pt thickness
        bottom.set(qn('w:space'), '0')
        bottom.set(qn('w:color'), 'auto')
        tblBorders.append(bottom)
        tblPr[0].append(tblBorders)


def add_metadata_surat(doc, context):
    # Zero-height anchor paragraph: without any paragraph between them, Word merges
    # this table into the Kop Surat table. It occupies ~1 pt, so no visible gap.
    anchor = doc.add_paragraph()
    anchor.paragraph_format.space_before = Pt(0)
    anchor.paragraph_format.space_after = Pt(0)
    anchor.paragraph_format.line_spacing = Pt(1)
    anchor.add_run().font.size = Pt(1)

    widths = [Cm(2.5), Cm(0.5), Cm(9.0), Cm(4.0)]  # total 16.0 cm
    table = doc.add_table(rows=3, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    # Borderless table
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        tblBorders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement(f'w:{edge}')
            el.set(qn('w:val'), 'nil')
            tblBorders.append(el)
        tblPr[0].append(tblBorders)

    for i, w in enumerate(widths):
        table.columns[i].width = w
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w

    data = [
        ("Nomor", context['nomor_surat'], context['tanggal_surat']),
        ("Lampiran", context['lampiran'], ""),
        ("Hal", context['hal'], ""),
    ]
    for row, (label, value, tanggal) in zip(table.rows, data):
        for cell, text in zip(row.cells, (label, ":", value, tanggal)):
            if text:
                cell.paragraphs[0].add_run(text)

    # Strict single spacing for every paragraph in every cell
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0

    # Date cell explicitly right-aligned
    cell_date = table.cell(0, 3)
    p_date = cell_date.paragraphs[0]
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.right_indent = Pt(0)

    # Ensure zero right cell margin on date cell so text aligns flush with the right edge
    tcPr = cell_date._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    right_mar = OxmlElement('w:right')
    right_mar.set(qn('w:w'), '0')
    right_mar.set(qn('w:type'), 'dxa')
    tcMar.append(right_mar)
    tcPr.append(tcMar)

    # Apply 4pt space before the first row to create a precise gap below the header line
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(4)


def add_dasar_penugasan(doc, dasar_text):
    # Paragraph "Berdasarkan:"
    p_dasar = doc.add_paragraph("Berdasarkan:")
    p_dasar.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_dasar.paragraph_format.first_line_indent = Cm(1.25)
    p_dasar.paragraph_format.space_after = Pt(0)
    p_dasar.paragraph_format.line_spacing = 1.5

    # Numbered List with Hanging Indent
    dasar_list = [d.strip() for d in dasar_text.split('\n') if d.strip()]
    for i, text in enumerate(dasar_list, start=1):
        p_item = doc.add_paragraph()
        p_item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_item.paragraph_format.left_indent = Cm(0.75)
        p_item.paragraph_format.first_line_indent = Cm(-0.75)
        p_item.paragraph_format.space_after = Pt(0)
        p_item.paragraph_format.line_spacing = 1.5
        p_item.add_run(f"{i}.\t{text}")


def add_body_paragraphs(doc, context):
    # Paragraph 1: kami menugasi (Flush left, NO indent)
    p1 = doc.add_paragraph(f"kami menugasi tim untuk melaksanakan {context['jenis_penugasan']} dengan susunan tim sesuai surat tugas terlampir.")
    p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p1.paragraph_format.space_after = Pt(0)
    p1.paragraph_format.line_spacing = 1.5

    # Paragraph 2: Penugasan ini dilaksanakan (First-line indent)
    p2 = doc.add_paragraph(f"Penugasan ini dilaksanakan selama {context['lama_hari']} hari kerja, terhitung mulai tanggal {context['tgl_mulai']} s.d. {context['tgl_selesai']}.")
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p2.paragraph_format.first_line_indent = Cm(1.25)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.5

    # Paragraph 3: Seluruh biaya (First-line indent)
    p3 = doc.add_paragraph(f"Seluruh biaya yang berkaitan dengan penugasan ini menjadi beban anggaran {context['beban_biaya']}.")
    p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p3.paragraph_format.first_line_indent = Cm(1.25)
    p3.paragraph_format.space_after = Pt(0)
    p3.paragraph_format.line_spacing = 1.5

    # Paragraph 4: Gratifikasi (First-line indent)
    p4 = doc.add_paragraph("Pegawai BPKP dalam melaksanakan tugas tidak menerima/meminta gratifikasi dan/atau suap.")
    p4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p4.paragraph_format.first_line_indent = Cm(1.25)
    p4.paragraph_format.space_after = Pt(0)
    p4.paragraph_format.line_spacing = 1.5

    # Paragraph 5: Penugasan ini agar dilaksanakan (First-line indent)
    p5 = doc.add_paragraph("Penugasan ini agar dilaksanakan dengan penuh tanggung jawab.")
    p5.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p5.paragraph_format.first_line_indent = Cm(1.25)
    p5.paragraph_format.space_after = Pt(0)
    p5.paragraph_format.line_spacing = 1.5


def add_closing_and_signature(doc, context):
    # Determine indent for the first line ONLY
    is_plt_plh = context['jabatan_ttd'].startswith("Plt.") or context['jabatan_ttd'].startswith("Plh.")
    jabatan_indent = Cm(9.25) if is_plt_plh else Cm(10.0)
    base_indent = Cm(10.0)  # Always 10.0 for everything below the title

    # 1-line gap before signature block
    p_gap_sig = doc.add_paragraph()
    p_gap_sig.paragraph_format.space_after = Pt(0)
    p_gap_sig.paragraph_format.space_before = Pt(0)
    p_gap_sig.paragraph_format.line_spacing = 1.0

    # Update helper to accept a dynamic indent value
    def set_signature_format(p, indent_val):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = indent_val
        p.paragraph_format.right_indent = Cm(0)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0

    # 1. Jabatan (Uses dynamic indent)
    p_jab = doc.add_paragraph(context['jabatan_ttd'] + ",")
    set_signature_format(p_jab, jabatan_indent)

    # Spacer (Uses base_indent)
    p_gap1 = doc.add_paragraph()
    set_signature_format(p_gap1, base_indent)

    # 2. TTE Spacer / Text (All use base_indent)
    if context.get('use_tte', True):
        p_tte = doc.add_paragraph()
        set_signature_format(p_tte, base_indent)
        run_tte = p_tte.add_run("Ditandatangani secara elektronik oleh")
        run_tte.font.name = 'Arial'
        run_tte.font.size = Pt(10)
        run_tte.font.italic = True
        run_tte.font.color.rgb = RGBColor(166, 166, 166)

        p_gap2 = doc.add_paragraph()
        set_signature_format(p_gap2, base_indent)
    else:
        p_gap2 = doc.add_paragraph()
        set_signature_format(p_gap2, base_indent)
        p_gap3 = doc.add_paragraph()
        set_signature_format(p_gap3, base_indent)

    # 3. Nama (Uses base_indent)
    p_nama = doc.add_paragraph(context['nama_ttd'])
    set_signature_format(p_nama, base_indent)

    # --- TEMBUSAN LOGIC ---
    # Parse items and auto-detect count
    raw_tembusan = context.get('daftar_tembusan') or context.get('tembusan') or ""
    items = [t.strip() for t in raw_tembusan.split('\n') if t.strip()]

    # Smart helper for Hanging Indent Tembusan Items (Multiple)
    def add_tembusan_items(parent_doc, tembusan_list):
        for i, text in enumerate(tembusan_list, start=1):
            # Strip existing punctuation to prevent double symbols
            clean_text = text.rstrip(';.')

            # Add ';' for items 1 to N-1, and '.' for the final item
            if i < len(tembusan_list):
                final_text = f"{clean_text};"
            else:
                final_text = f"{clean_text}."

            p = parent_doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.add_run(f"{i}.\t{final_text}")

    posisi = context.get('posisi_tembusan', 'Halaman 1 (Langsung)')

    if len(items) == 1:
        # RULE 8a: Single Item (Auto-detect)
        p_tembusan = doc.add_paragraph("Tembusan:")
        p_tembusan.paragraph_format.left_indent = Cm(0)
        p_tembusan.paragraph_format.first_line_indent = Cm(0)
        p_tembusan.paragraph_format.space_before = Pt(12)
        p_tembusan.paragraph_format.space_after = Pt(0)
        p_tembusan.paragraph_format.line_spacing = 1.0

        # No numbers, add "Yth.", no punctuation at the end
        p_item = doc.add_paragraph(f"Yth. {items[0].rstrip(';.')}")
        p_item.paragraph_format.left_indent = Cm(0)
        p_item.paragraph_format.first_line_indent = Cm(0)
        p_item.paragraph_format.space_after = Pt(0)
        p_item.paragraph_format.line_spacing = 1.0

    elif len(items) > 1:
        if posisi == "Halaman 1 (Langsung)":
            # RULE 8b: Multiple Items on Page 1
            p_tembusan = doc.add_paragraph("Tembusan Yth.:")
            p_tembusan.paragraph_format.left_indent = Cm(0)
            p_tembusan.paragraph_format.first_line_indent = Cm(0)
            p_tembusan.paragraph_format.space_before = Pt(12)
            p_tembusan.paragraph_format.space_after = Pt(0)
            p_tembusan.paragraph_format.line_spacing = 1.0

            add_tembusan_items(doc, items)

        elif posisi == "Halaman 2 (Daftar Terlampir)":
            # RULE 8c: Multiple Items on Page 2 (Header only on Page 1)
            p_tembusan = doc.add_paragraph("Tembusan Yth. (Daftar Terlampir)")
            p_tembusan.paragraph_format.left_indent = Cm(0)
            p_tembusan.paragraph_format.first_line_indent = Cm(0)
            p_tembusan.paragraph_format.space_before = Pt(12)
            p_tembusan.paragraph_format.space_after = Pt(0)
            p_tembusan.paragraph_format.line_spacing = 1.0

    # --- HOTLINE LOGIC ---
    p_hotline = doc.add_paragraph()
    p_hotline.paragraph_format.left_indent = Cm(0)
    p_hotline.paragraph_format.first_line_indent = Cm(0)
    p_hotline.paragraph_format.space_before = Pt(12)
    p_hotline.paragraph_format.space_after = Pt(0)
    p_hotline.paragraph_format.line_spacing = 1.0

    run_1 = p_hotline.add_run("Layanan Pengaduan/")
    run_1.font.name = 'Arial'
    run_1.font.size = Pt(10)

    run_hotline = p_hotline.add_run("Hotline")
    run_hotline.font.name = 'Arial'
    run_hotline.font.size = Pt(10)
    run_hotline.italic = True

    run_2 = p_hotline.add_run(":\nApabila Pegawai BPKP menerima/meminta gratifikasi dan/atau suap, dapat dilaporkan melalui mekanisme penyampaian pengaduan pada wbs.bpkp.go.id. atau menghubungi 081389101133.")
    run_2.font.name = 'Arial'
    run_2.font.size = Pt(10)

    # --- PAGE 2 LAMPIRAN LOGIC (VIA HEADER) ---
    if len(items) > 1 and posisi == "Halaman 2 (Daftar Terlampir)":
        # Add a new section (this automatically creates a Next Page Section Break)
        new_section = doc.add_section()

        # Adjust Margins specifically for Section 2
        new_section.header_distance = Cm(0.5)  # Pushes Header up to 0.5cm from top edge
        new_section.top_margin = Cm(2.5)       # Pulls Body text up, removing huge gaps

        # Unlink the header so Page 1 remains completely clean
        header = new_section.header
        header.is_linked_to_previous = False

        # Write to the Header with a shifted left indent (7.0 cm) to prevent text wrapping
        p_lamp = header.paragraphs[0] if len(header.paragraphs) > 0 else header.add_paragraph()
        p_lamp.text = "" 
        p_lamp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_lamp.paragraph_format.left_indent = Cm(7.0)  # Shifted left to 7.0cm
        p_lamp.paragraph_format.right_indent = Cm(0)
        p_lamp.paragraph_format.first_line_indent = Cm(0)
        p_lamp.paragraph_format.space_before = Pt(0)
        p_lamp.paragraph_format.space_after = Pt(0)
        p_lamp.paragraph_format.line_spacing = 1.0

        run1 = p_lamp.add_run("Lampiran\n")
        run2 = p_lamp.add_run("Surat Kepala Perwakilan BPKP Provinsi Papua\n")
        run3 = p_lamp.add_run(f"Nomor: {context['nomor_surat']}\n")
        run4 = p_lamp.add_run(f"Tanggal: {context['tanggal_surat']}")

        # Force font size 11 for the header
        for r in p_lamp.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(11)

        # Add exactly 1 blank line with 1.0 spacing before the title
        p_gap_judul = doc.add_paragraph()
        p_gap_judul.paragraph_format.space_before = Pt(0)
        p_gap_judul.paragraph_format.space_after = Pt(0)
        p_gap_judul.paragraph_format.line_spacing = 1.0

        # Write to the Body of Section 2
        p_judul = doc.add_paragraph("DAFTAR TEMBUSAN")
        p_judul.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_judul.paragraph_format.space_before = Pt(0) 
        p_judul.paragraph_format.space_after = Pt(12)
        p_judul.runs[0].bold = True

        add_tembusan_items(doc, items)


def add_tembusan(doc, context):
    """Fallback helper for tembusan; formatting is integrated in add_closing_and_signature."""
    pass


def generate_sp2_doc(context):
    doc = docx.Document()
    set_document_margins(doc)

    # Reuse perfected Kop and Metadata
    add_kop_surat(doc)
    add_metadata_surat(doc, context)

    # Add Gap before Yth (Strict 1.0 spacing, 0pt before/after)
    p_gap_yth = doc.add_paragraph()
    p_gap_yth.paragraph_format.space_after = Pt(0)
    p_gap_yth.paragraph_format.space_before = Pt(0)
    p_gap_yth.paragraph_format.line_spacing = 1.0

    # Add Yth.
    p_yth = doc.add_paragraph(f"Yth. {context['yth']}")
    p_yth.paragraph_format.left_indent = Cm(0)
    p_yth.paragraph_format.space_after = Pt(0)
    p_yth.paragraph_format.line_spacing = 1.0

    # Add di Tempat
    p_di = doc.add_paragraph(f"di {context['tempat_tujuan']}")
    p_di.paragraph_format.left_indent = Cm(0)
    p_di.paragraph_format.space_after = Pt(0)
    p_di.paragraph_format.line_spacing = 1.0

    # Add Gap before body
    p_gap_body = doc.add_paragraph()
    p_gap_body.paragraph_format.space_after = Pt(0)
    p_gap_body.paragraph_format.space_before = Pt(0)
    p_gap_body.paragraph_format.line_spacing = 1.0

    # Paragraf 1
    teks_paragraf_1 = (
        f"Berdasarkan hasil pemantauan atas pelaksanaan tindak lanjut rekomendasi Laporan "
        f"Nomor: {context['nomor_laporan']} tanggal {context['tanggal_laporan']} hal "
        f"{context['hal_laporan']}, masih terdapat rekomendasi pada Organisasi Perangkat "
        f"Daerah (OPD) terkait di lingkungan {context['nama_pemda']} yang belum kami terima "
        f"dokumen/bukti tindak lanjutnya, sebagaimana terlampir."
    )
    p_par1 = doc.add_paragraph(teks_paragraf_1)
    p_par1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_par1.paragraph_format.left_indent = Cm(0)
    p_par1.paragraph_format.right_indent = Cm(0)
    p_par1.paragraph_format.first_line_indent = Cm(1.25)
    p_par1.paragraph_format.space_before = Pt(0)
    p_par1.paragraph_format.space_after = Pt(0)
    p_par1.paragraph_format.line_spacing = 1.5

    # Paragraf 2
    teks_paragraf_2 = (
        f"Sehubungan dengan hal tersebut, kami mengharapkan {context['sapaan_tujuan']} "
        f"melakukan percepatan pelaksanaan tindak lanjut dimaksud. Bukti tindak lanjut "
        f"disampaikan kepada Kepala Perwakilan BPKP Provinsi Papua pada kesempatan pertama. "
        f"Untuk koordinasi lebih lanjut dapat menghubungi {context['sapaan_cp']} "
        f"{context['nama_cp']} (HP: {context['no_hp_cp']})."
    )
    p_par2 = doc.add_paragraph(teks_paragraf_2)
    p_par2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_par2.paragraph_format.left_indent = Cm(0)
    p_par2.paragraph_format.right_indent = Cm(0)
    p_par2.paragraph_format.first_line_indent = Cm(1.25)
    p_par2.paragraph_format.space_before = Pt(0)
    p_par2.paragraph_format.space_after = Pt(0)
    p_par2.paragraph_format.line_spacing = 1.5

    # Paragraf 3 (Penutup)
    teks_paragraf_3 = (
        f"Atas perhatian dan kerja sama {context['sapaan_tujuan']}, kami mengucapkan terima kasih."
    )
    p_par3 = doc.add_paragraph(teks_paragraf_3)
    p_par3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_par3.paragraph_format.left_indent = Cm(0)
    p_par3.paragraph_format.right_indent = Cm(0)
    p_par3.paragraph_format.first_line_indent = Cm(1.25)
    p_par3.paragraph_format.space_before = Pt(0)
    p_par3.paragraph_format.space_after = Pt(0)
    p_par3.paragraph_format.line_spacing = 1.5

    # Re-use Signature and Tembusan modules from Surat Pengantar
    add_closing_and_signature(doc, context)

    if context.get('tembusan') and str(context['tembusan']).strip():
        add_tembusan(doc, context)

    # --- STEP 2: LAMPIRAN 1 (FORM TP3) ---
    from docx.enum.section import WD_SECTION

    # Create a new section for Lampiran 1 to have a separate header
    section_lamp1 = doc.add_section(WD_SECTION.NEW_PAGE)
    section_lamp1.header_distance = Cm(0.5)

    # Access and unlink the header so it doesn't affect the main letter
    header_lamp1 = section_lamp1.header
    header_lamp1.is_linked_to_previous = False

    # Clear any default empty paragraphs in the new header
    for hp in header_lamp1.paragraphs:
        hp.text = ""

    # Build the Lampiran Header text
    p_lamp1 = header_lamp1.paragraphs[0] if header_lamp1.paragraphs else header_lamp1.add_paragraph()
    p_lamp1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_lamp1.paragraph_format.left_indent = Cm(7.0)
    p_lamp1.paragraph_format.space_before = Pt(0)
    p_lamp1.paragraph_format.space_after = Pt(0)
    p_lamp1.paragraph_format.line_spacing = 1.0

    run_lamp1 = p_lamp1.add_run(
        f"Lampiran 1\n"
        f"Surat Kepala Perwakilan BPKP Provinsi Papua\n"
        f"Nomor: {context['nomor_surat']}\n"
        f"Tanggal: {context['tanggal_surat']}"
    )
    run_lamp1.font.name = 'Arial'
    run_lamp1.font.size = Pt(12)

    # Process Form TP3 if uploaded (extract ONLY page 1) or leave blank placeholder
    if context.get('file_tp3') is not None:
        import fitz  # PyMuPDF
        import io

        file_tp3 = context['file_tp3']
        if hasattr(file_tp3, 'seek'):
            file_tp3.seek(0)
        pdf_bytes = file_tp3.read() if hasattr(file_tp3, 'read') else file_tp3.getvalue()
        pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")

        if len(pdf_doc) > 0:
            # ONLY extract and insert the first page (page 1)
            page = pdf_doc.load_page(0)
            pix = page.get_pixmap(dpi=200)
            img_bytes = pix.tobytes("png")
            img_stream = io.BytesIO(img_bytes)

            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(0)
            p_img.paragraph_format.space_after = Pt(0)

            run_img = p_img.add_run()
            run_img.add_picture(img_stream, width=Cm(16.0))
        else:
            p_blank = doc.add_paragraph()
            p_blank.paragraph_format.space_before = Pt(0)
            p_blank.paragraph_format.space_after = Pt(0)
    else:
        # If no file uploaded, leave the rest of the page blank as a placeholder
        p_blank = doc.add_paragraph()
        p_blank.paragraph_format.space_before = Pt(0)
        p_blank.paragraph_format.space_after = Pt(0)

    # --- STEP 3: LAMPIRAN 2 (TABEL REKOMENDASI LANDSCAPE) ---
    from docx.enum.section import WD_SECTION, WD_ORIENT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn, nsdecls

    # 1. Create Landscape Section
    section_lamp2 = doc.add_section(WD_SECTION.NEW_PAGE)
    section_lamp2.orientation = WD_ORIENT.LANDSCAPE

    # Swap dimensions for A4 Landscape
    new_width, new_height = section_lamp2.page_height, section_lamp2.page_width
    section_lamp2.page_width = new_width
    section_lamp2.page_height = new_height
    section_lamp2.header_distance = Cm(0.5)

    # 1b. Force page numbering to restart at 1 for this section
    sectPr = section_lamp2._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:start'), '1')

    # 2. Setup Independent Header
    header_lamp2 = section_lamp2.header
    header_lamp2.is_linked_to_previous = False
    for hp in header_lamp2.paragraphs:
        hp.text = ""

    p_head2 = header_lamp2.paragraphs[0] if header_lamp2.paragraphs else header_lamp2.add_paragraph()
    p_head2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_head2.paragraph_format.left_indent = Cm(15.5)
    p_head2.paragraph_format.space_before = Pt(0)
    p_head2.paragraph_format.space_after = Pt(0)
    p_head2.paragraph_format.line_spacing = 1.0

    # Helper function to inject Word dynamic field codes
    def add_field(run, field_type):
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = field_type
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run._r.append(fldChar3)

    # Build Header Text with dynamic numbering: "Lampiran 2 / [PAGE] - [SECTIONPAGES]"
    run_h1 = p_head2.add_run("Lampiran 2 / ")
    run_h1.font.name = 'Arial'
    run_h1.font.size = Pt(12)

    run_page = p_head2.add_run()
    run_page.font.name = 'Arial'
    run_page.font.size = Pt(12)
    add_field(run_page, 'PAGE')

    run_h2 = p_head2.add_run(" - ")
    run_h2.font.name = 'Arial'
    run_h2.font.size = Pt(12)

    run_secpages = p_head2.add_run()
    run_secpages.font.name = 'Arial'
    run_secpages.font.size = Pt(12)
    add_field(run_secpages, 'SECTIONPAGES')

    run_h3 = p_head2.add_run(
        f"\nSurat Kepala Perwakilan BPKP Provinsi Papua\n"
        f"Nomor: {context['nomor_surat']}\n"
        f"Tanggal: {context['tanggal_surat']}\n"
    )
    run_h3.font.name = 'Arial'
    run_h3.font.size = Pt(12)

    # 3. Add Document Title (Immediately after header setup, no spacer paragraph)
    p_title = doc.add_paragraph("DAFTAR REKOMENDASI YANG BELUM DITINDAKLANJUTI")
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.runs[0]
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(12)
    run_title.bold = True

    # Prepare Data from DataFrame
    import pandas as pd
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml

    df = context.get('df_rekom')

    # Filter out rows that are entirely empty or just whitespace securely
    if df is not None and not df.empty:
        # Replace empty strings or whitespace-only strings with None, then drop rows where all columns are None
        df = df.replace(r'^\s*$', None, regex=True).dropna(how='all')

    num_data_rows = len(df) if (df is not None and not df.empty) else 4
    total_rows = num_data_rows + 1

    # 4. Create the Table
    table = doc.add_table(rows=total_rows, cols=4)
    table.style = 'Table Grid'

    widths = [Cm(1.5), Cm(7.5), Cm(7.5), Cm(8.2)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

    # 5. Format Header Row & Set Repeat Header XML
    headers = ["No.", "Kondisi", "Penyebab", "Rekomendasi yang Belum Ditindaklanjuti"]
    for idx, text in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = text

        shading_elm = parse_xml(r'<w:shd {} w:fill="D9D9D9"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading_elm)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        run = p.runs[0]
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.bold = True

    # Inject XML to Repeat Header Row
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

    # 6. Format Body Rows and Inject Data
    for r_idx in range(1, total_rows):
        # Extract data securely
        if df is not None and not df.empty and (r_idx - 1) < len(df):
            row_data = df.iloc[r_idx - 1]
            val_kondisi = str(row_data.get('Kondisi', '')) if pd.notna(row_data.get('Kondisi')) else ""
            val_penyebab = str(row_data.get('Penyebab', '')) if pd.notna(row_data.get('Penyebab')) else ""
            val_rekom = str(row_data.get('Rekomendasi yang Belum Ditindaklanjuti', '')) if pd.notna(row_data.get('Rekomendasi yang Belum Ditindaklanjuti')) else ""
        else:
            val_kondisi, val_penyebab, val_rekom = "", "", ""

        cols_text = [f"{r_idx}.", val_kondisi, val_penyebab, val_rekom]

        for c_idx in range(4):
            cell = table.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(cols_text[c_idx])
            run.font.name = 'Arial'
            run.font.size = Pt(12)

    return doc


def render_surat_pengantar():
    st.title("Drafting Surat Pengantar")

    # Build the Form
    with st.form("tnd_form"):
        st.subheader("Data Naskah Dinas")
        col1, col2 = st.columns(2)
        with col1:
            nomor_surat = st.text_input("Nomor Surat", value="")
            lampiran = st.text_input("Lampiran", value="")
            hal = st.text_input("Hal", value="")
            tanggal_surat_input = st.date_input("Tanggal Dokumen", datetime.date.today())
            is_draft = st.checkbox("Simpan sebagai Draft (Tanggal Kosong)", value=False)

        with col2:
            yth = st.text_input("Yth.", value="")
            tempat_tujuan = st.text_input("Kota Tujuan (di)", value="")

        st.subheader("Isi Surat")
        ada_dasar_lain = st.checkbox("Tambahkan Dasar Aturan Lain (selain PP 60 & Perpres 2)", value=False)
        dasar_aturan = st.text_area("Isi Dasar Aturan Tambahan (jika dicentang)", value="")

        jenis_penugasan = st.text_input("kami menugasi tim untuk melaksanakan:", value="")

        col3, col4, col5 = st.columns(3)
        with col3:
            lama_hari = st.text_input("Lama Hari Kerja", value="", placeholder="misal: 36 (tiga puluh enam)")
        with col4:
            tgl_mulai_input = st.date_input("Mulai Tanggal", datetime.date.today())
        with col5:
            tgl_selesai_input = st.date_input("S.d. Tanggal", datetime.date.today())

        beban_biaya = st.text_input("Beban Anggaran", value="")

        st.subheader("Penandatangan")
        col6, col7 = st.columns(2)
        with col6:
            jabatan_ttd = st.text_input("Jabatan", value="")
            nama_ttd = st.text_input("Nama", value="")
        with col7:
            checker_tte = st.checkbox("Gunakan Tanda Tangan Elektronik (TTE)", value=True)
            checker_plt = st.selectbox("Status Jabatan", ["Definitif", "Plt.", "Plh."])

        st.subheader("Daftar Tembusan")
        posisi_tembusan = st.radio("Posisi Tembusan:", ["Halaman 1 (Langsung)", "Halaman 2 (Daftar Terlampir)"])
        daftar_tembusan = st.text_area("Isi Tembusan (satu baris per nama)")

        submit_button = st.form_submit_button("Generate Draft Surat")

    # Logic when form is submitted
    if submit_button:
        if not nomor_surat or not hal or not nama_ttd:
            st.error("⚠️ Harap isi form yang kosong terlebih dahulu!")
        else:
            # Update Context dynamically
            jabatan_raw = f"{checker_plt + ' ' if checker_plt != 'Definitif' else ''}{jabatan_ttd}"

            context = {
                'nomor_surat': nomor_surat,
                'tanggal_surat': format_tgl_indo(tanggal_surat_input, is_draft),
                'lampiran': lampiran,
                'hal': hal,
                'yth': yth,
                'tempat_tujuan': tempat_tujuan,
                'jenis_penugasan': jenis_penugasan,
                'lama_hari': lama_hari,
                'tgl_mulai': format_tgl_indo(tgl_mulai_input, include_year=(tgl_mulai_input.year != tgl_selesai_input.year)),
                'tgl_selesai': format_tgl_indo(tgl_selesai_input),
                'beban_biaya': beban_biaya,
                'jabatan_ttd': jabatan_raw.strip(),  # STRIP to prevent center-alignment shifting
                'nama_ttd': nama_ttd.strip(),
                'use_tte': checker_tte,  # Capture the boolean from the checkbox
            }
            context['posisi_tembusan'] = posisi_tembusan
            context['daftar_tembusan'] = daftar_tembusan

            # Handle Dasar Instruksi conditionally with STRICT comma at the end
            dasar_tambahan = dasar_aturan.strip().rstrip(';.,')

            if ada_dasar_lain and dasar_tambahan:
                dasar_instruksi_final = (
                    "Peraturan Pemerintah Nomor 60 Tahun 2008 tentang Sistem Pengendalian Intern Pemerintah;\n"
                    "Peraturan Presiden Nomor 2 Tahun 2025 tentang Perubahan Kedua atas Peraturan Presiden Nomor 192 Tahun 2014 tentang Badan Pengawasan Keuangan dan Pembangunan; dan\n"
                    f"{dasar_tambahan},"
                )
            else:
                dasar_instruksi_final = (
                    "Peraturan Pemerintah Nomor 60 Tahun 2008 tentang Sistem Pengendalian Intern Pemerintah; dan\n"
                    "Peraturan Presiden Nomor 2 Tahun 2025 tentang Perubahan Kedua atas Peraturan Presiden Nomor 192 Tahun 2014 tentang Badan Pengawasan Keuangan dan Pembangunan,"
                )

            doc = docx.Document()
            set_document_margins(doc)

            add_kop_surat(doc)
            add_metadata_surat(doc, context)

            # Add Gap before Yth
            p_gap_yth = doc.add_paragraph()
            p_gap_yth.paragraph_format.space_after = Pt(0)
            p_gap_yth.paragraph_format.space_before = Pt(0)
            p_gap_yth.paragraph_format.line_spacing = 1.0

            # Add Yth.
            p_yth = doc.add_paragraph(f"Yth. {context['yth']}")
            p_yth.paragraph_format.left_indent = Cm(0)
            p_yth.paragraph_format.space_after = Pt(0)
            p_yth.paragraph_format.line_spacing = 1.0

            # Add di Tempat
            p_di = doc.add_paragraph(f"di {context['tempat_tujuan']}")
            p_di.paragraph_format.left_indent = Cm(0)
            p_di.paragraph_format.space_after = Pt(0)
            p_di.paragraph_format.line_spacing = 1.0

            # Gap before Berdasarkan
            p_gap_dasar = doc.add_paragraph()
            p_gap_dasar.paragraph_format.space_after = Pt(0)
            p_gap_dasar.paragraph_format.space_before = Pt(0)
            p_gap_dasar.paragraph_format.line_spacing = 1.0

            add_dasar_penugasan(doc, dasar_instruksi_final)
            add_body_paragraphs(doc, context)
            add_closing_and_signature(doc, context)

            preview_path = os.path.join(os.getcwd(), "preview_sementara.docx")
            doc.save(preview_path)

            # Streamlit Download Button instead of os.startfile
            with open(preview_path, "rb") as docx_file:
                st.download_button(
                    label="📥 Download Naskah Dinas",
                    data=docx_file,
                    file_name="Surat_Pengantar_TND.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

            st.success("✅ Naskah Dinas berhasil dibuat! Silakan klik tombol di atas untuk mengunduh.")


def render_sp2():
    st.title("Generator Surat Penegasan (SP2)")

    # 1. MAIN METADATA FORM
    with st.form("sp2_form"):
        st.subheader("Data Naskah Dinas (SP2)")
        col1, col2 = st.columns(2)
        with col1:
            nomor_surat = st.text_input("Nomor Surat", value="")
            lampiran = st.text_input("Lampiran", value="")
            tanggal_surat_input = st.date_input("Tanggal Dokumen", datetime.date.today())
            is_draft = st.checkbox("Simpan sebagai Draft (Tanggal Kosong)", value=False, key="draft_sp2")

        with col2:
            yth = st.text_input("Yth.", value="")
            tempat_tujuan = st.text_input("Kota Tujuan (di)", value="")

        st.subheader("Data Laporan Sebelumnya (Paragraf 1)")
        col3, col4 = st.columns(2)
        with col3:
            nomor_laporan = st.text_input("Nomor Laporan", value="")
            tanggal_laporan_input = st.date_input("Tanggal Laporan", datetime.date.today(), key="tgl_lap")
        with col4:
            hal_laporan = st.text_input("Hal Laporan", value="")
            nama_pemda = st.text_input("Nama Pemda/Lingkungan", value="")

        st.subheader("Data Tindak Lanjut & Kontak (Paragraf 2)")
        col5, col6 = st.columns(2)
        with col5:
            sapaan_tujuan = st.selectbox("Sapaan Tujuan", ["Bapak", "Ibu"])
            sapaan_cp = st.selectbox("Sapaan Contact Person", ["Sdr.", "Sdri."])
        with col6:
            nama_cp = st.text_input("Nama Contact Person", value="")
            no_hp_cp = st.text_input("Nomor HP CP", value="")

        st.subheader("Penandatangan & Tembusan")
        col7, col8 = st.columns(2)
        with col7:
            jabatan_ttd = st.text_input("Jabatan Penandatangan", value="")
            nama_ttd = st.text_input("Nama Penandatangan", value="")
            tembusan = st.text_area("Tembusan (pisahkan per baris dengan Enter)", value="")
        with col8:
            checker_tte = st.checkbox("Gunakan Tanda Tangan Elektronik (TTE)", value=True, key="tte_sp2")

        st.subheader("Upload Form TP3 (Untuk Lampiran 1)")
        file_tp3 = st.file_uploader("Upload PDF Form TP3 (Opsional)", type=["pdf"], key="tp3_upload")

        submit_sp2 = st.form_submit_button("Generate SP2")

    # 2. LAMPIRAN 2 UI (OUTSIDE & BELOW FORM)
    st.subheader("Data Rekomendasi (Lampiran 2)")
    metode_lampiran2 = st.radio("Metode Pengisian Tabel:", ["Input Langsung via Web", "Generate Tabel Kosong (Isi manual di Word)"])

    df_to_pass = None
    if metode_lampiran2 == "Input Langsung via Web":
        st.caption("Ketik langsung atau Paste dari Excel. Kolom 'No.' akan terisi otomatis di dokumen.")
        import pandas as pd
        if "df_rekom" not in st.session_state:
            st.session_state.df_rekom = pd.DataFrame(
                columns=["Kondisi", "Penyebab", "Rekomendasi yang Belum Ditindaklanjuti"],
                data=[["", "", ""]]
            )

        df_to_pass = st.data_editor(
            st.session_state.df_rekom,
            num_rows="dynamic",
            use_container_width=True,
            hide_index=True,
            key="data_editor_rekom"
        )

    if submit_sp2:
        # Mapping to standard context dictionary
        context = {
            'nomor_surat': nomor_surat,
            'lampiran': lampiran,
            'hal': "Penegasan Kedua atas Rekomendasi Pengawasan yang Belum Ditindaklanjuti",
            'tanggal_surat': format_tgl_indo(tanggal_surat_input, is_draft),
            'yth': yth,
            'tempat_tujuan': tempat_tujuan,
            'nomor_laporan': nomor_laporan,
            'tanggal_laporan': format_tgl_indo(tanggal_laporan_input, False),
            'hal_laporan': hal_laporan,
            'nama_pemda': nama_pemda,
            'sapaan_tujuan': sapaan_tujuan,
            'sapaan_cp': sapaan_cp,
            'nama_cp': nama_cp,
            'no_hp_cp': no_hp_cp,
            'jabatan_ttd': jabatan_ttd,
            'nama_ttd': nama_ttd,
            'use_tte': checker_tte,
            'tembusan': tembusan,
            'file_tp3': file_tp3,
            'df_rekom': df_to_pass
        }

        # Generate the document
        doc = generate_sp2_doc(context)

        import tempfile
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            preview_path = tmp.name
            doc.save(preview_path)

        with open(preview_path, "rb") as docx_file:
            st.download_button(
                label="📥 Download Draft SP2",
                data=docx_file,
                file_name="Draft_SP2.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        st.success("✅ Fondasi SP2 berhasil dibuat!")


st.sidebar.title("Navigasi TND")
modul_pilihan = st.sidebar.selectbox("Pilih Modul Naskah Dinas:", ["Surat Pengantar", "Surat Penegasan (SP2)"])

if modul_pilihan == "Surat Pengantar":
    render_surat_pengantar()
elif modul_pilihan == "Surat Penegasan (SP2)":
    render_sp2()

